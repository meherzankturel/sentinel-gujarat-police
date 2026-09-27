#!/usr/bin/env python3
"""
A .pptx carrying the same deck, for a portal that will not take a PDF.

The slides are rendered images rather than native text boxes, deliberately.
Rebuilding this design in PowerPoint shapes would produce a worse-looking
deck that also disagreed with the PDF within a week. The PDF is the
primary; this exists so a file-type restriction cannot stop a submission.

Each slide carries its text as speaker notes, so the deck is still usable
to present from and its content is still searchable.

    python tools/build_pptx.py        # needs docs/deck/Sentinel-Presentation.pdf
"""
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile

ROOT = pathlib.Path(__file__).resolve().parent.parent
DECK = ROOT / "docs" / "deck"
PDF = DECK / "Sentinel-Presentation.pdf"
OUT = DECK / "Sentinel-Presentation.pptx"
TEMPLATE = DECK / "template.html"


def slide_notes() -> list[str]:
    """Visible text per slide, recovered from the template."""
    html = TEMPLATE.read_text()
    notes = []
    for block in re.findall(r"<section class=\"slide.*?</section>", html, re.S):
        txt = re.sub(r"<script.*?</script>", " ", block, flags=re.S)
        txt = re.sub(r"<[^>]+>", " ", txt)
        txt = txt.replace("&amp;", "&").replace("&nbsp;", " ")
        txt = re.sub(r"\s+", " ", txt).strip()
        notes.append(txt)
    return notes


def main() -> int:
    if not PDF.exists():
        sys.exit("run tools/build_deck.py --pdf first")
    if not shutil.which("pdftoppm"):
        sys.exit("pdftoppm not found (brew install poppler)")

    from pptx import Presentation
    from pptx.util import Emu

    work = pathlib.Path(tempfile.mkdtemp(prefix="sentinel-pptx-"))
    subprocess.run(["pdftoppm", "-png", "-r", "144", str(PDF), str(work / "s")],
                   check=True, capture_output=True)
    pages = sorted(work.glob("s-*.png"))
    if not pages:
        sys.exit("pdftoppm produced no pages")

    prs = Presentation()
    # 13.333in x 7.5in -- the 16:9 size the PDF was laid out at.
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]

    notes = slide_notes()
    for i, png in enumerate(pages):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), 0, 0,
                                 width=prs.slide_width,
                                 height=prs.slide_height)
        if i < len(notes):
            slide.notes_slide.notes_text_frame.text = notes[i]

    prs.save(OUT)
    shutil.rmtree(work, ignore_errors=True)
    print(f"wrote {OUT.relative_to(ROOT)}  "
          f"({len(pages)} slides, {OUT.stat().st_size/1_000_000:.1f} MB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
